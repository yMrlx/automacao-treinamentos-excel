"""Gera planilhas de DEMONSTRAÇÃO (dados 100% fictícios) no layout que o programa espera.

Uso (da raiz do repositório):
    python demo/gerar_planilhas_demo.py

Cria, ao lado de ``_Automação/``:

    Planos e Macs/
        Planos e Macs Agosto 2026.xlsx                      <- mês anterior (ponto de partida)
        POPs do Sistema/
            POPs aprovados e efetivos - Setembro 2026.xlsx  <- export do sistema de documentos
            POPs Obsoletos - Setembro 2026.xlsx

Com esses arquivos, o ciclo de Setembro/2026 (``python _Automação/main.py sync
--mes 9 --ano 2026 --vigentes "<export>"``, ou o painel) mostra tudo o que o
programa faz: POP que versionou com prazo correndo (ON TIME), POP que versionou
há tempo (ATRASADO), POP que sumiu do export (VERIFICAR), status preservados (OK
de quem elaborou o POP, HSE, NA) e o resumo no rodapé de cada aba.

Precisa de ``openpyxl`` (está no ``requirements-dev.txt``).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill

RAIZ = Path(__file__).resolve().parents[1]
PASTA = RAIZ / "Planos e Macs"
PASTA_POPS = PASTA / "POPs do Sistema"

VERDE = PatternFill("solid", fgColor="FF00B050")
CINZA = PatternFill("solid", fgColor="FFD9D9D9")
CABECALHO = PatternFill("solid", fgColor="FF305496")
NEGRITO_BRANCO = Font(bold=True, color="FFFFFFFF")
NEGRITO = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

# --- Catálogo fictício de POPs do time ------------------------------------------
# código: (título, versão no Agosto, versão no export de Setembro, aprovação em Set)
POPS = {
    "DOC-POP-0001001": ("Controle de acesso lógico", "3.0", "3.0", datetime(2026, 3, 10)),
    "DOC-POP-0001002": ("Gestão de mudanças em sistemas", "4.0", "5.0", datetime(2026, 9, 17)),
    "DOC-POP-0001003": ("Backup e restauração de dados", "2.0", "2.0", datetime(2025, 11, 4)),
    "DOC-POP-0001004": ("Gestão de incidentes de TI", "6.0", "6.0", datetime(2026, 1, 22)),
    "DOC-POP-0001005": ("Validação de sistemas computadorizados", "2.0", "3.0", datetime(2026, 9, 22)),
    "DOC-POP-0001006": ("Integridade de dados (ALCOA+)", "1.0", "1.0", datetime(2025, 6, 30)),
    "DOC-POP-0001007": ("Gestão de senhas e credenciais", "6.0", "7.0", datetime(2026, 8, 20)),
    "DOC-POP-0001008": ("Gestão de fornecedores de TI", "5.0", "5.0", datetime(2026, 2, 14)),
    # Este some do export de Setembro (foi obsoletado) -> a Matriz mostra VERIFICAR.
    "DOC-POP-0001009": ("Inventário de ativos de TI", "2.0", None, datetime(2025, 9, 1)),
    "DOC-POP-0001010": ("Plano de continuidade de negócio", "8.0", "8.0", datetime(2026, 4, 3)),
}

# POPs de outras áreas, só pra o export parecer o de verdade (ninguém do time tem).
OUTROS_POPS = [
    (f"DOC-POP-00020{n:02d}", f"Procedimento de {area}", f"{(n % 7) + 1}.0", datetime(2026, (n % 9) + 1, (n % 27) + 1))
    for n, area in enumerate(
        ["Produção", "Qualidade", "Almoxarifado", "Manutenção", "RH", "Financeiro",
         "Laboratório", "Expedição", "Compras", "Jurídico", "Marketing", "Engenharia"],
        start=1,
    )
]

# --- Pessoas fictícias e seus status por aba de cargo ------------------------------
# aba: (rótulo do cargo, [(id, nome)], blocos=[[códigos]], status[(pessoa, código)])
ABAS = {
    "Analista": ("Analista de Sistemas", [(101, "Ana Ribeiro"), (102, "Bruno Costa"), (103, "Carla Mendes")],
                 [["DOC-POP-0001001", "DOC-POP-0001002", "DOC-POP-0001007"],
                  ["DOC-POP-0001005", "DOC-POP-0001006", "DOC-POP-0001009"]]),
    "Analista Sr": ("Analista de Sistemas Sênior", [(201, "Daniel Souza"), (202, "Elisa Martins")],
                    [["DOC-POP-0001002", "DOC-POP-0001004", "DOC-POP-0001005"]]),
    "Administrativo": ("Analista ADM", [(301, "Fernanda Lima")],
                       [["DOC-POP-0001003", "DOC-POP-0001007", "DOC-POP-0001008"]]),
    "Gerente": ("Gerente IS", [(401, "Gustavo Rocha")],
                [["DOC-POP-0001002", "DOC-POP-0001007", "DOC-POP-0001010"]]),
    "Terceiros": ("Terceiros", [(501, "Hugo Almeida"), (502, "Igor Santos")],
                  [["DOC-POP-0001001", "DOC-POP-0001006"]]),
    "Estagiario": ("Estagiário", [(601, "Júlia Pereira")],
                   [["DOC-POP-0001001", "DOC-POP-0001002", "DOC-POP-0001005"]]),
}

# Exceções ao "OK" padrão: (aba, pessoa, código) -> (planejado, status)
EXCECOES = {
    ("Analista", "Bruno Costa", "DOC-POP-0001002"): ("ELABORADOR", "OK"),  # quem escreveu o POP
    ("Analista", "Carla Mendes", "DOC-POP-0001006"): ("05/09/2026", "ON TIME"),  # vence sem versionar
    ("Analista Sr", "Elisa Martins", "DOC-POP-0001004"): ("15/07/2026", "ATRASADO"),
    ("Administrativo", "Fernanda Lima", "DOC-POP-0001008"): (None, "NA"),
    ("Gerente", "Gustavo Rocha", "DOC-POP-0001010"): (None, "HSE"),
    ("Terceiros", "Igor Santos", "DOC-POP-0001001"): (None, "GLOBAL"),
}


def _largura(ws, colunas: dict[str, float]) -> None:
    for letra, largura in colunas.items():
        ws.column_dimensions[letra].width = largura


def _aba_versionamento(ws, versoes: dict[str, tuple]) -> None:
    cab = ["Document Number", "Title", "Status", "Version", "Owning Department",
           "Impacted Departments", "Approved Date", "Approver", "Owner", "Reviewer"]
    ws.append(cab)
    for celula in ws[1]:
        celula.fill, celula.font = CABECALHO, NEGRITO_BRANCO
    for codigo, (titulo, versao, aprovado) in versoes.items():
        ws.append([codigo, titulo, "Em vigor", versao, "TI", "TI", aprovado,
                   "Aprovador Fictício", "Dono Fictício", "Revisor Fictício"])
    for linha in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        linha[0].number_format = "dd/mm/yyyy"
    _largura(ws, {"A": 18, "B": 42, "C": 10, "D": 9, "G": 13})


def _matriz(ws) -> None:
    ws["C3"], ws["D3"], ws["E3"] = "Códigos dos POPs (abaixo)", "Descrição", "Versão em que o time treinou"
    ws["C4"], ws["D4"], ws["E4"], ws["F4"] = "Código", "Descrição do Procedimento", "Versão", "Versionamento"
    for celula in ("C4", "D4", "E4", "F4"):
        ws[celula].fill, ws[celula].font, ws[celula].alignment = CABECALHO, NEGRITO_BRANCO, CENTRO
    for i, (codigo, (titulo, versao_ago, _set, _aprov)) in enumerate(POPS.items(), start=5):
        ws[f"C{i}"] = codigo
        ws[f"D{i}"] = titulo
        ws[f"E{i}"] = versao_ago  # TEXTO, como a operadora digita
        ws[f"F{i}"] = f"=IFERROR(VLOOKUP(C{i},'Versionamento Mês'!$A:$D,4,0),\"VERIFICAR\")"
    ultima = 4 + len(POPS)
    # Mesmo alerta visual da planilha real: F diferente de E fica vermelho.
    ws.conditional_formatting.add(
        f"F5:F{ultima}",
        CellIsRule(operator="notEqual", formula=["E5"], fill=PatternFill("solid", fgColor="FFFF0000")),
    )
    _largura(ws, {"C": 18, "D": 40, "E": 12, "F": 15})


def _aba_de_cargo(ws, rotulo: str, pessoas: list, blocos: list) -> None:
    ws["A1"] = "Plano de Treinamento"
    ws["A1"].font = Font(bold=True, size=14)
    linha = 3
    # Bloco de integração SEM código de POP (existe na planilha real): o programa
    # não o versiona, mas ele delimita o bloco de cima (fim de bloco).
    blocos_desenho = [("integracao", ["Integração de Qualidade", "Integração de HSE"])] + [
        ("pops", codigos) for codigos in blocos
    ]
    for tipo, itens in blocos_desenho:
        ws.cell(linha, 1, " Treinamentos").font = NEGRITO
        ws.cell(linha + 1, 1, "Códigos dos Treinamentos").font = NEGRITO
        ws.cell(linha + 2, 1, "Versão").font = NEGRITO
        ws.cell(linha + 3, 1, "ID").font = NEGRITO
        ws.cell(linha + 3, 2, rotulo).font = NEGRITO
        for j, item in enumerate(itens):
            col = 3 + 2 * j
            if tipo == "pops":
                ws.cell(linha, col, f"=VLOOKUP({ws.cell(linha + 1, col).coordinate},'Matriz - Atualização'!$C:$D,2,0)")
                ws.cell(linha + 1, col, item)
                ws.cell(linha + 2, col, f"=VLOOKUP({ws.cell(linha + 1, col).coordinate},'Matriz - Atualização'!$C:$E,3,0)")
            else:
                ws.cell(linha, col, item)
                ws.cell(linha + 1, col, "Integração")
                ws.cell(linha + 2, col, "N/A")
            for k in (col, col + 1):
                for r in (linha, linha + 1, linha + 2):
                    ws.cell(r, k).fill = CINZA
                    ws.cell(r, k).alignment = CENTRO
            ws.cell(linha + 3, col, "Planejado").font = NEGRITO
            ws.cell(linha + 3, col + 1, "Executado").font = NEGRITO
            for p, (pid, nome) in enumerate(pessoas):
                r = linha + 4 + p
                ws.cell(r, 1, pid)
                ws.cell(r, 2, nome)
                if tipo == "pops":
                    planejado, status = EXCECOES.get((ws.title, nome, item), (None, "OK"))
                else:
                    planejado, status = None, "Integração"
                p_cel = ws.cell(r, col, planejado)
                if isinstance(planejado, str):
                    p_cel.number_format = "@"  # data como TEXTO, como a operadora digita
                c = ws.cell(r, col + 1, status)
                c.alignment = CENTRO
                if status == "OK":
                    c.fill = VERDE
        linha += 4 + len(pessoas) + 2
    # Rodapé do mês anterior: registro de auditoria que se acumula mês a mês.
    linha += 1
    ws.cell(linha, 1, "Referente Agosto 2026").font = NEGRITO
    ws.cell(linha + 1, 1, "Não houveram alterações nos POPs")
    _largura(ws, {"A": 26, "B": 22, **{chr(ord("C") + i): 14 for i in range(8)}})


def gerar() -> None:
    PASTA_POPS.mkdir(parents=True, exist_ok=True)

    # 1) Planos e Macs de Agosto (o mês anterior, de onde o ciclo parte).
    wb = Workbook()
    _aba_versionamento(
        wb.active,
        {c: (t, v, a if v_set == v else datetime(2026, 7, 10))
         for c, (t, v, v_set, a) in POPS.items()},
    )
    wb.active.title = "Versionamento Mês"
    _matriz(wb.create_sheet("Matriz - Atualização"))
    for aba, (rotulo, pessoas, blocos) in ABAS.items():
        _aba_de_cargo(wb.create_sheet(aba), rotulo, pessoas, blocos)
    wb.save(PASTA / "Planos e Macs Agosto 2026.xlsx")

    # 2) Export de POPs vigentes de Setembro (o sistema de documentos).
    vigentes = {c: (t, v_set, a) for c, (t, _v, v_set, a) in POPS.items() if v_set}
    vigentes.update({c: (t, v, a) for c, t, v, a in OUTROS_POPS})
    wb = Workbook()
    wb.active.title = "Export"
    _aba_versionamento(wb.active, vigentes)
    wb.save(PASTA_POPS / "POPs aprovados e efetivos - Setembro 2026.xlsx")

    # 3) Export de POPs obsoletos de Setembro.
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(["Document Number", "Title", "Status", "Major Version Number", "Obsolete Date"])
    ws.append(["DOC-POP-0001009", "Inventário de ativos de TI", "Obsoleto", "2", datetime(2026, 9, 12)])
    ws.append(["DOC-POP-0000777", "Procedimento antigo de impressão", "Obsoleto", "3", datetime(2026, 9, 3)])
    wb.save(PASTA_POPS / "POPs Obsoletos - Setembro 2026.xlsx")

    for arquivo in sorted(PASTA.rglob("*.xlsx")):
        print("gerado:", arquivo.relative_to(RAIZ))


if __name__ == "__main__":
    gerar()
